"""
excel.py

Manejo del Excel Maestro y procesamiento del reporte diario de alarmas.
Basado strictly en día calendario (00:00 a 23:59).
"""

from datetime import datetime, timedelta
from pathlib import Path
from typing import Union, Tuple
import re
import os
import tempfile
import shutil
import unicodedata
from datetime import date
import openpyxl
import pandas as pd


class ExcelManager:

    EQUIVALENCIAS_PARQUES = {
        "EL PELICANO": "El Pelicano",
        "RAUQUEN": "CL114 Rauquen",
        "CHILLAN 1": "CL121 Chillan I",
        "CHILLAN I": "CL121 Chillan I",
        "EL PASO": "CL126 El Paso",
        "MOLINA": "CL130 Molina",
        "VILLA ALEGRE": "CL125 Villa Alegre",
        "MUTUPIN": "CL129 Mutupin",
        "CHILLAN 2": "CL122 Chillan II",
        "CHILLAN II": "CL122 Chillan II",
        "VILLA CRUZ": "CL109 Villa Cruz",
        "PEGASUS": "CL124 Pegasus",
        "LEMU": "CL113 Lemu",
        "CARENA": "CL213 Carena",
        "ORION": "CL123 Orion",
        "PENCAHUE": "CL205 Pencahue",
        "SANTA CAROLINA": "CL132 Santa Carolina",
        "SANTA FE": "CL120 Santa Fe",
        "BATRES": "CL204 Batres",
        "PEQUEN": "CL209 Pequen",
        "LINARES": "CL127 Linares",
        "PACHIRA": "CL128 Pachira",
        "EL ROMERAL": "CL131 El Romeral",
        "BLB": "CL210 BLB",
        "VILLA SECA": "CL118 Villa Seca",
    }

    NOMBRE_HOJA_PLANTILLA = "PLANTILLA"

    # Días de la semana en inglés para coincidir con el diseño de la tabla
    DIAS_INGLES = ["Friday", "Saturday", "Sunday", "Monday", "Tuesday", "Wednesday", "Thursday"]

    def __init__(self, archivo_maestro: Union[str, Path]):
        self.archivo_maestro = Path(archivo_maestro) if archivo_maestro else None

    def existe(self) -> bool:
        if not self.archivo_maestro:
            return False
        return self.archivo_maestro.exists()

    def _calcular_turno_y_fecha(self, dt: datetime):
        fecha = dt.date()
        hora = dt.time()
        t_inicio = datetime.strptime("08:00:00", "%H:%M:%S").time()
        t_fin = datetime.strptime("20:00:00", "%H:%M:%S").time()

        if t_inicio <= hora < t_fin:
            return fecha, "DAY"
        else:
            return fecha, "NIGHT"

    @staticmethod
    def _normalizar_nombre(valor):
        texto = unicodedata.normalize("NFKD", str(valor or ""))
        texto = "".join(c for c in texto if not unicodedata.combining(c))
        return " ".join(texto.upper().split())

    @staticmethod
    def _inicio_semana(fecha):
        return fecha - timedelta(days=(fecha.weekday() - 4) % 7)

    def procesar_reporte_alarmas(self, archivo_nuevo: Path):
        """Una única semana operacional (viernes-jueves) por archivo.

        Rechaza fechas inválidas y reportes que mezclan semanas: nunca descarta
        registros de forma silenciosa.
        """
        df = pd.read_excel(archivo_nuevo, sheet_name="Registro de alarmas y eventos", header=8)
        col_hora = "Hora de activación (cliente)"
        col_region = "Región"
        for col in (col_hora, col_region):
            if col not in df.columns:
                raise ValueError(f"Falta columna requerida en el reporte: {col}")
        df = df.loc[df[col_hora].notna()].copy()
        if df.empty:
            raise ValueError("El reporte no contiene alarmas fechadas. No se modificó el Maestro.")
        df["dt"] = pd.to_datetime(df[col_hora], errors="coerce", dayfirst=True)
        invalidas = df["dt"].isna()
        if invalidas.any():
            raise ValueError(f"Hay {invalidas.sum()} fechas inválidas; no se modificó el Maestro.")
        # Los eventos NVR sin región no bloquean las alarmas válidas.
        # Se contabilizan como excepciones; jamás se asignan a un parque arbitrario.
        sin_parque = df[col_region].isna() | df[col_region].astype(str).str.strip().eq("")
        self.ultimo_sin_region = int(sin_parque.sum())
        self.ultimo_total_leido = len(df)
        equivalencias = {self._normalizar_nombre(k): v for k, v in self.EQUIVALENCIAS_PARQUES.items()}
        nombres_oficiales = {self._normalizar_nombre(v): v for v in self.EQUIVALENCIAS_PARQUES.values()}
        def mapear(valor):
            clave = self._normalizar_nombre(valor)
            return equivalencias.get(clave, nombres_oficiales.get(clave, str(valor).strip()))
        df["Parque_Oficial"] = df[col_region].apply(mapear)
        df.loc[sin_parque, "Parque_Oficial"] = None
        fechas_turnos = df["dt"].apply(self._calcular_turno_y_fecha)
        df["Fecha_Reporte"] = [x[0] for x in fechas_turnos]
        df["Turno"] = [x[1] for x in fechas_turnos]
        semanas = {self._inicio_semana(f) for f in df["Fecha_Reporte"]}
        if len(semanas) != 1:
            rangos = ", ".join(str(f) for f in sorted(semanas))
            raise ValueError("El archivo contiene más de una semana operacional (viernes-jueves): " + rangos + ". Sepáralas antes de importar; no se modificó el Maestro.")
        viernes_inicio = semanas.pop()
        # Identificador ISO correspondiente al lunes de la semana del viernes inicial.
        numero_semana = (viernes_inicio - timedelta(days=4)).isocalendar().week
        conteo = df.loc[~sin_parque].groupby(["Fecha_Reporte", "Parque_Oficial", "Turno"]).size().unstack(fill_value=0)
        for col in ("DAY", "NIGHT"):
            if col not in conteo.columns:
                conteo[col] = 0
        return conteo.reset_index(), numero_semana, viernes_inicio

    @staticmethod
    def _rango_hoja(ws):
        """Extrae fechas reales de A1; evita identificar semanas solo por WW."""
        valor = ws.cell(1, 1).value
        if not isinstance(valor, str):
            return None
        coincidencia = re.search(r"(\d{2}-\d{2}-\d{4})\s+to\s+(\d{2}-\d{2}-\d{4})", valor)
        if not coincidencia:
            return None
        try:
            return (datetime.strptime(coincidencia.group(1), "%d-%m-%Y").date(),
                    datetime.strptime(coincidencia.group(2), "%d-%m-%Y").date())
        except ValueError:
            return None

    def _preparar_encabezados_fechas(self, ws, viernes_inicio: datetime.date):
        """Escribe el rango semanal en A1 y formatea las columnas B-O en Fila 2

        siguiendo el estilo de la plantilla: 'Alarms on [Día] [Número Día] Day / Night'
        """
        jueves_cierre = viernes_inicio + timedelta(days=6)

        # 1. Rango de fecha en la Fila 1 (ej: 07-08-2026 to 13-08-2026)
        ws.cell(
            row=1,
            column=1,
            value=f"{viernes_inicio.strftime('%d-%m-%Y')} to {jueves_cierre.strftime('%d-%m-%Y')}",
        )

        # 2. Encabezados de días (Fila 2) desde Columna B (2) hasta O (15)
        col_actual = 2
        for i in range(7):
            fecha_dia = viernes_inicio + timedelta(days=i)
            nombre_dia = self.DIAS_INGLES[i]

            # Formato: "Alarms on Friday 14 Day / Night"
            texto_encabezado = f"Alarms on {nombre_dia} {fecha_dia.day} Day / Night"

            ws.cell(row=2, column=col_actual, value=texto_encabezado)
            col_actual += 2

    def agregar_registros(self, archivo_nuevo: Path, log_callback=None) -> dict:
        def log(mensaje):
            if log_callback:
                log_callback(mensaje)

        df_resumen, numero_semana, viernes_inicio = self.procesar_reporte_alarmas(archivo_nuevo)
        jueves = viernes_inicio + timedelta(days=6)
        wb = openpyxl.load_workbook(self.archivo_maestro)
        log(f"EXCEL: {self.ultimo_total_leido} eventos leídos; {self.ultimo_sin_region} sin región (excluidos y reportados).")
        try:
            # Identificar la hoja correspondiente, incluso si A1 estaba mal históricamente.
            hojas_fecha = [ws for ws in wb.worksheets
                           if ws.title != self.NOMBRE_HOJA_PLANTILLA
                           and self._rango_hoja(ws) == (viernes_inicio, jueves)]
            if len(hojas_fecha) > 1:
                raise ValueError("Hay varias hojas para el mismo período; revisión manual requerida.")
            nombre = f"Monitoreo WW {numero_semana:02d}"
            if hojas_fecha:
                ws_destino = hojas_fecha[0]
            elif nombre in wb.sheetnames:
                ws_destino = wb[nombre]
                rango_existente = self._rango_hoja(ws_destino)
                if rango_existente is not None and rango_existente != (viernes_inicio, jueves):
                    # Si A1 es incorrecto, verificar las fechas del encabezado B2:O2.
                    encabezado = str(ws_destino.cell(2, 2).value or "")
                    if not re.search(rf"\bFriday\s+{viernes_inicio.day}\b", encabezado, re.I):
                        raise ValueError(f"'{nombre}' tiene otro período y encabezados incompatibles; no se modificó.")
            else:
                ws_destino = None
            if self.NOMBRE_HOJA_PLANTILLA not in wb.sheetnames:
                raise ValueError("Falta la hoja PLANTILLA. No se creará una hoja desde datos históricos.")
            hoja_base = wb[self.NOMBRE_HOJA_PLANTILLA]
            parques = {}
            fila_total = None
            primera = ultima = None
            for fila in range(3, hoja_base.max_row + 1):
                valor = hoja_base.cell(fila, 1).value
                if valor is None:
                    continue
                clave = self._normalizar_nombre(valor)
                if clave == "TOTAL":
                    fila_total = fila
                elif clave not in ("PHOTOVOLTAIC PARK", "FECHA", "TURNO", ""):
                    if clave in parques:
                        raise ValueError(f"Parque duplicado en PLANTILLA: {valor}")
                    parques[clave] = fila
                    primera = fila if primera is None else primera
                    ultima = fila
            if not parques:
                raise ValueError("La PLANTILLA no tiene parques en la columna A.")
            presentes = {self._normalizar_nombre(p) for p in df_resumen["Parque_Oficial"]}
            desconocidos = sorted(presentes - set(parques))
            # La plantilla define los únicos parques contabilizables.
            excluidos_parque = df_resumen[df_resumen["Parque_Oficial"].map(self._normalizar_nombre).isin(desconocidos)]
            total_excluidos_parque = int(excluidos_parque[["DAY", "NIGHT"]].sum().sum())
            if desconocidos:
                log("EXCEL: Parques no incluidos en Maestro (excluidos): " + ", ".join(desconocidos))
            df_resumen = df_resumen[~df_resumen["Parque_Oficial"].map(self._normalizar_nombre).isin(desconocidos)].copy()
            total_omitidos = self.ultimo_sin_region + total_excluidos_parque
            log(f"EXCEL: {total_omitidos} eventos excluidos ({self.ultimo_sin_region} sin región, {total_excluidos_parque} de parques fuera de plantilla).")

            if ws_destino is None:
                ws = wb.copy_worksheet(hoja_base)
                ws.title = nombre
                creado = True
            else:
                ws = ws_destino
                creado = False
                log(f"EXCEL: Actualizando semana existente '{ws.title}' (sin sumar duplicados).")
            self._preparar_encabezados_fechas(ws, viernes_inicio)
            # Reemplazar los conteos de la semana importada, incluso cuando son cero.
            # Solo tocar las filas de parques existentes en la plantilla.
            for fila in parques.values():
                for col in range(2, 16):
                    ws.cell(fila, col, 0)
            for registro in df_resumen.itertuples(index=False):
                dia = (registro.Fecha_Reporte - viernes_inicio).days
                if not 0 <= dia <= 6:
                    raise ValueError(f"Fecha fuera de semana: {registro.Fecha_Reporte}")
                fila = parques[self._normalizar_nombre(registro.Parque_Oficial)]
                ws.cell(fila, 2 + dia * 2, int(registro.DAY))
                ws.cell(fila, 3 + dia * 2, int(registro.NIGHT))
            if fila_total is not None and primera is not None and ultima is not None:
                for col in range(2, 16):
                    letra = openpyxl.utils.get_column_letter(col)
                    ws.cell(fila_total, col, f"=SUM({letra}{primera}:{letra}{ultima})")
            total = int(df_resumen[["DAY", "NIGHT"]].sum().sum())
            # Escritura atómica: el Maestro original no se toca hasta terminar.
            maestro = Path(self.archivo_maestro)
            fd, temporal = tempfile.mkstemp(prefix=".weeklyreport_", suffix=maestro.suffix, dir=maestro.parent)
            os.close(fd)
            try:
                wb.save(temporal)
                comprobacion = openpyxl.load_workbook(temporal, read_only=True)
                try:
                    if ws.title not in comprobacion.sheetnames:
                        raise ValueError("No se pudo verificar la hoja creada.")
                finally:
                    comprobacion.close()
                # Copia de respaldo antes de reemplazar el archivo original.
                respaldo = maestro.with_name(maestro.stem + "_backup_" + datetime.now().strftime("%Y%m%d_%H%M%S_%f") + maestro.suffix)
                shutil.copy2(maestro, respaldo)
                os.replace(temporal, maestro)
            finally:
                if os.path.exists(temporal):
                    os.unlink(temporal)
            log(f"EXCEL: {'Creada' if creado else 'Actualizada'} '{ws.title}' ({viernes_inicio:%d/%m/%Y} al {jueves:%d/%m/%Y}); {total} alarmas. Respaldo: {respaldo.name}")
            return {"total": self.ultimo_total_leido, "agregados": total, "omitidos": total_omitidos, "existente": not creado}
        finally:
            wb.close()
